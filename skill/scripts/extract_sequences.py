#!/usr/bin/env python3
"""
专利生物序列导出器
通过智慧芽 Bio OpenAPI 提取单篇专利中的生物序列，支持导出为 Excel / FASTA 格式。

Usage:
    python extract_sequences.py \
        --api-key <KEY> \
        --patent-number WO2004050828A2 \
        --seq-id-range "1-50" \
        --output-format excel \
        --output-dir /path/to/output
"""

import argparse
import sys
import time
from datetime import datetime
from pathlib import Path

import requests
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, PatternFill

# ── API 配置 ──────────────────────────────────────────────
API_URL = "https://connect.zhihuiya.com/bio-openapi/patent/extract/single"
PAGE_SIZE = 100
MAX_RETRIES = 3          # 请求失败最大重试次数（含首次）
RETRY_INTERVAL = 3       # 重试间隔（秒）


def _timestamped_path(path: str) -> str:
    """文件被占用时生成带时间戳的新路径。"""
    ts = datetime.now().strftime("%Y%m%d%H%M%S")
    p = Path(path)
    return str(p.parent / f"{p.stem}_{ts}{p.suffix}")


# ── Seq ID 解析 ────────────────────────────────────────────
def parse_seq_id_range(seq_input: str) -> tuple:
    """
    根据用户输入推断 extract_type 和 sequence_no_range。

    - "all" / "全部"           → EXTRACT_ALL, "all"
    - 含 "-" (如 "1-50")       → EXTRACT_RANGE
    - 纯逗号分隔 (如 "1,5,10") → EXTRACT_NUM
    - 混合 (如 "1-10,23,24")   → EXTRACT_RANGE
    """
    raw = seq_input.strip()
    lower = raw.lower()

    if lower in ("all", "全部", "all sequences", "所有", "全部序列"):
        return "EXTRACT_ALL", "all"

    has_range = "-" in raw
    if has_range:
        return "EXTRACT_RANGE", raw
    else:
        return "EXTRACT_NUM", raw


# ── API 调用 ───────────────────────────────────────────────
def fetch_page(api_key, patent_number, patent_id, extract_type,
               sequence_no_range, sequence_type, sequence_length,
               page, size=PAGE_SIZE):
    """调用 API 获取单页序列数据（带重试）。"""
    headers = {
        "Authorization": f"Bearer {api_key}",
        "Content-Type": "application/json;charset=UTF-8",
    }

    payload = {
        "extract_type": extract_type,
        "page": page,
        "size": size,
        "sequence_no_range": sequence_no_range,
    }

    if patent_id:
        payload["patent_id"] = patent_id
    elif patent_number:
        payload["patent_number"] = patent_number
    else:
        raise ValueError("patent_number 和 patent_id 至少提供一个")

    if sequence_type:
        payload["sequence_type"] = sequence_type
    if sequence_length:
        payload["sequence_length"] = sequence_length

    last_err = None
    for attempt in range(1, MAX_RETRIES + 1):
        try:
            resp = requests.post(API_URL, headers=headers, json=payload, timeout=120)

            if resp.status_code != 200:
                raise RuntimeError(
                    f"API 请求失败 HTTP {resp.status_code}: {resp.text[:500]}"
                )

            body = resp.json()
            if not body.get("status", False):
                err_code = body.get("error_code", "unknown")
                err_msg = body.get("error_msg") or body.get("error_message") or ""
                raise RuntimeError(f"API 返回错误 error_code={err_code}, error_msg={err_msg}")

            return body

        except Exception as e:
            last_err = e
            if attempt < MAX_RETRIES:
                print(f"[warning] 第 {attempt} 次请求失败: {e}，{RETRY_INTERVAL}s 后重试...", file=sys.stderr)
                time.sleep(RETRY_INTERVAL)
            else:
                print(f"[error] 已重试 {MAX_RETRIES} 次仍失败: {e}", file=sys.stderr)

    raise RuntimeError(f"API 请求在 {MAX_RETRIES} 次重试后仍失败: {last_err}")


def fetch_all(api_key, patent_number, patent_id, extract_type,
              sequence_no_range, sequence_type, sequence_length):
    """分页拉取全部序列。"""
    all_results = []
    total = None
    page = 1

    while True:
        body = fetch_page(
            api_key, patent_number, patent_id, extract_type,
            sequence_no_range, sequence_type, sequence_length,
            page=page,
        )
        data = body.get("data") or {}
        if total is None:
            total = data.get("total", 0)
            print(f"[info] 序列总数: {total}", file=sys.stderr)

        results = data.get("results") or []
        all_results.extend(results)

        # 分页进度百分比
        pct = 100.0 * len(all_results) / total if total else 100.0
        print(
            f"[info] 第 {page} 页: 获取 {len(results)} 条，"
            f"累计 {len(all_results)}/{total} ({pct:.1f}%)",
            file=sys.stderr,
        )

        if not results or len(all_results) >= total:
            break
        page += 1

    return all_results, total


# ── Excel 导出 ─────────────────────────────────────────────
def export_excel(sequences, output_path, patent_label):
    """导出为 Excel，含 Summary 工作表 + Sequences 工作表。location 对象展平为多列。"""
    wb = Workbook()

    # ── Summary 工作表 ──
    ws_summary = wb.active
    ws_summary.title = "Summary"

    # 序列类型分布统计
    type_dist = {}
    for s in sequences:
        st = s.get("sequence_type", "UNKNOWN")
        type_dist[st] = type_dist.get(st, 0) + 1

    summary_rows = [
        ("专利号", patent_label),
        ("序列总数", len(sequences)),
        ("导出时间", datetime.now().strftime("%Y-%m-%d %H:%M:%S")),
        ("", ""),
        ("序列类型分布", ""),
    ]
    for st, cnt in sorted(type_dist.items()):
        summary_rows.append((st, cnt))

    # 写 Summary
    s_font = Font(bold=True, color="FFFFFF")
    s_fill = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")

    for r, (k, v) in enumerate(summary_rows, 1):
        c1 = ws_summary.cell(row=r, column=1, value=k)
        c2 = ws_summary.cell(row=r, column=2, value=v)
        if r == 1 or (k and not v and r in (5,)):
            c1.font = s_font
            c1.fill = s_fill
        elif k and not isinstance(v, (int, float)) and not v:
            c1.font = Font(bold=True)

    ws_summary.column_dimensions["A"].width = 20
    ws_summary.column_dimensions["B"].width = 35

    # ── Sequences 工作表 ──
    ws = wb.create_sheet("Sequences")

    # 收集所有 location 键
    loc_keys = set()
    for s in sequences:
        loc = s.get("location") or {}
        loc_keys.update(loc.keys())
    loc_keys = sorted(loc_keys)

    headers = (
        ["Seq ID", "Sequence ID Number", "Sequence Type", "Length"]
        + [f"Location_{k}" for k in loc_keys]
        + ["Sequence"]
    )

    # 写表头
    h_font = Font(bold=True, color="FFFFFF")
    h_fill = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")
    h_align = Alignment(horizontal="center", vertical="center", wrap_text=True)

    for col, h in enumerate(headers, 1):
        c = ws.cell(row=1, column=col, value=h)
        c.font = h_font
        c.fill = h_fill
        c.alignment = h_align

    # 写数据行
    for r, s in enumerate(sequences, 2):
        sid = s.get("seq_id", "")
        sid_num = s.get("sequence_id_number", [])
        sid_num_str = ", ".join(str(x) for x in sid_num) if isinstance(sid_num, list) else str(sid_num)
        stype = s.get("sequence_type", "")
        slen = s.get("s_len", "")
        seq = s.get("s_sequence", "")
        loc = s.get("location") or {}

        col = 1
        ws.cell(row=r, column=col, value=sid); col += 1
        ws.cell(row=r, column=col, value=sid_num_str); col += 1
        ws.cell(row=r, column=col, value=stype); col += 1
        ws.cell(row=r, column=col, value=slen); col += 1
        for lk in loc_keys:
            lv = loc.get(lk, [])
            lv_str = ", ".join(str(x) for x in lv) if isinstance(lv, list) else str(lv)
            ws.cell(row=r, column=col, value=lv_str); col += 1

        cell = ws.cell(row=r, column=col, value=seq)
        cell.alignment = Alignment(wrap_text=False)

    # 列宽
    for ci in range(1, len(headers) + 1):
        max_len = len(str(headers[ci - 1]))
        for ri in range(2, min(len(sequences) + 2, 50)):
            v = ws.cell(row=ri, column=ci).value
            if v is not None:
                max_len = max(max_len, min(len(str(v)), 40))
        ws.column_dimensions[ws.cell(row=1, column=ci).column_letter].width = max_len + 2

    ws.freeze_panes = "A2"

    try:
        wb.save(output_path)
    except PermissionError:
        output_path = _timestamped_path(output_path)
        wb.save(output_path)
        print(f"[warning] 文件被占用，已改用: {output_path}", file=sys.stderr)
    print(f"[done] Excel 导出: {output_path} ({len(sequences)} 条序列)")


# ── FASTA 导出 ─────────────────────────────────────────────
def _write_fasta_inner(sequences, f, patent_label):
    """写入 FASTA 内容到已打开的文件对象。"""
    for s in sequences:
        sid = s.get("seq_id", "")
        sid_num = s.get("sequence_id_number", [])
        sid_num_str = ",".join(str(x) for x in sid_num) if isinstance(sid_num, list) else str(sid_num)
        stype = s.get("sequence_type", "UNKNOWN")
        slen = s.get("s_len", 0)
        seq = s.get("s_sequence", "")

        header = f">{patent_label}|{sid_num_str}|{stype}|{slen}|seq_id:{sid}"
        f.write(header + "\n")

        for i in range(0, len(seq), 60):
            f.write(seq[i:i + 60] + "\n")


def export_fasta(sequences, output_path, patent_label):
    """导出为 FASTA，header: >专利号|SEQ_ID_NO|类型|长度|seq_id:ID"""
    try:
        with open(output_path, "w") as f:
            _write_fasta_inner(sequences, f, patent_label)
    except PermissionError:
        output_path = _timestamped_path(output_path)
        with open(output_path, "w") as f:
            _write_fasta_inner(sequences, f, patent_label)
        print(f"[warning] 文件被占用，已改用: {output_path}", file=sys.stderr)
    print(f"[done] FASTA 导出: {output_path} ({len(sequences)} 条序列)")


# ── 主入口 ─────────────────────────────────────────────────
def main():
    parser = argparse.ArgumentParser(
        description="通过智慧芽 Bio OpenAPI 提取专利生物序列并导出 Excel/FASTA"
    )
    parser.add_argument("--api-key", required=True, help="API 密钥")
    parser.add_argument("--patent-number", help="专利公开号 (如 WO2004050828A2)")
    parser.add_argument("--patent-id", help="专利 UUID (优先于 patent-number)")
    parser.add_argument("--seq-id-range", required=True,
                        help='Seq ID 范围，如 "1-50"、"1,5,10"、"1-10,23,24"、"all"')
    parser.add_argument("--output-format", choices=["excel", "fasta", "both"],
                        default="excel", help="输出格式 (默认 excel)")
    parser.add_argument("--output-dir", required=True, help="输出目录")
    parser.add_argument("--sequence-type", choices=["PROTEIN", "NUCLEOTIDE"],
                        help="序列类型过滤 (可选)")
    parser.add_argument("--seq-length-start", type=int, help="序列长度范围起始 (可选)")
    parser.add_argument("--seq-length-end", type=int, help="序列长度范围结束 (可选)")
    args = parser.parse_args()

    # 校验专利标识
    if not args.patent_number and not args.patent_id:
        print("[error] 必须提供 --patent-number 或 --patent-id", file=sys.stderr)
        sys.exit(1)

    patent_label = args.patent_id or args.patent_number

    # 推断 extract_type
    extract_type, seq_no_range = parse_seq_id_range(args.seq_id_range)
    print(f"[info] 专利: {patent_label}", file=sys.stderr)
    print(f"[info] extract_type={extract_type}, sequence_no_range={seq_no_range}",
          file=sys.stderr)

    # 构造 sequence_length
    seq_len = None
    if args.seq_length_start is not None or args.seq_length_end is not None:
        seq_len = {
            "start": args.seq_length_start if args.seq_length_start is not None else 0,
            "end": args.seq_length_end if args.seq_length_end is not None else 999999,
        }

    # 拉取序列
    try:
        sequences, total = fetch_all(
            api_key=args.api_key,
            patent_number=args.patent_number,
            patent_id=args.patent_id,
            extract_type=extract_type,
            sequence_no_range=seq_no_range,
            sequence_type=args.sequence_type,
            sequence_length=seq_len,
        )
    except Exception as e:
        print(f"[error] 拉取序列失败: {e}", file=sys.stderr)
        sys.exit(1)

    if not sequences:
        print("[warning] 未找到符合条件的序列", file=sys.stderr)
        sys.exit(0)

    print(f"\n[info] 共获取 {len(sequences)} 条序列 (API 报告总数: {total})",
          file=sys.stderr)

    # 导出
    out_dir = Path(args.output_dir)
    out_dir.mkdir(parents=True, exist_ok=True)

    if args.output_format in ("excel", "both"):
        p = out_dir / f"{patent_label}_sequences.xlsx"
        export_excel(sequences, str(p), patent_label)

    if args.output_format in ("fasta", "both"):
        p = out_dir / f"{patent_label}_sequences.fasta"
        export_fasta(sequences, str(p), patent_label)


if __name__ == "__main__":
    main()
