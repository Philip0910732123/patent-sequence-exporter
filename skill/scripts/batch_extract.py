#!/usr/bin/env python3
"""
专利生物序列批量提取工具 - Skill 版本 v1.0
从智慧芽 Bio OpenAPI 批量提取专利中的生物序列。

与 exe 版本的区别：
- 支持多专利批量提取（exe 仅单专利）
- NDJSON 即时落盘 + 断点恢复
- 进度回调（供 AI 平台实时展示）
- 非交互式（多结果取第一个，exe 需用户手动选择）
- CLI 入口，无 Flask/Web UI
"""

import argparse
import json
import os
import re
import sys
import time
from datetime import datetime
from pathlib import Path

import requests

try:
    from openpyxl import Workbook
    from openpyxl.styles import Font, Alignment, PatternFill
    HAS_OPENPYXL = True
except ImportError:
    HAS_OPENPYXL = False

# ── 配置 ──────────────────────────────────────────────────
API_URL = "https://connect.zhihuiya.com/bio-openapi/patent/extract/single"
PATENT_SEARCH_URL = "https://connect.zhihuiya.com/search/patent/patent-search-pn"
PAGE_SIZE = 100
MAX_RETRIES = 3
RETRY_INTERVAL = 3
BATCH_SIZE = 10
BATCH_INTERVAL = 5


# ── 工具函数 ──────────────────────────────────────────────

def parse_patent_ids(raw_input):
    """解析专利号列表，支持逗号/分号/换行分隔。"""
    parts = re.split(r'[,;\n]+', raw_input.strip())
    result = []
    seen = set()
    for p in parts:
        p = p.strip().replace(' ', '').replace('/', '')
        if p and p not in seen:
            seen.add(p)
            result.append(p)
    return result


def parse_seq_id_range(seq_input):
    """解析 Seq ID 范围输入。"""
    raw = seq_input.strip()
    if ' ' in raw and ',' not in raw and '-' not in raw:
        parts = raw.split()
        raw = ','.join(parts)
    lower = raw.lower()
    if lower in ('all', '全部', 'all sequences', '所有', '全部序列'):
        return 'EXTRACT_ALL', 'all'
    if '-' in raw:
        return 'EXTRACT_RANGE', raw
    return 'EXTRACT_NUM', raw


def read_api_key(api_key, api_key_file=None):
    """从直接输入或文件路径获取 API Key。"""
    if api_key:
        return api_key.strip()
    if api_key_file:
        p = Path(api_key_file)
        if p.exists():
            content = p.read_text(encoding='utf-8-sig').strip()
            m = re.search(r'sk-[A-Za-z0-9]+', content)
            if m:
                return m.group(0)
            for line in content.splitlines():
                line = line.strip()
                if line and not line.startswith('#') and not line.startswith('//'):
                    if ':' in line:
                        val = line.split(':', 1)[1].strip()
                        if val and val.isascii():
                            return val
                    if '=' in line:
                        val = line.split('=', 1)[1].strip()
                        if val and val.isascii():
                            return val
                    if line.isascii() and len(line) > 10:
                        return line
            raise ValueError('无法从文件中提取 API Key')
        raise ValueError(f'文件不存在: {api_key_file}')
    raise ValueError('未提供 API Key')


def _timestamped_path(path):
    """生成带时间戳的文件路径，避免文件占用冲突。"""
    ts = datetime.now().strftime('%Y%m%d%H%M%S')
    p = Path(path)
    return str(p.parent / f'{p.stem}_{ts}{p.suffix}')


# ── 专利号标准化 API ─────────────────────────────────────

def search_patent_pn(api_key, query, authority=None):
    """调用专利号标准化 API，返回标准化专利号列表。"""
    headers = {
        'Authorization': f'Bearer {api_key}',
        'Content-Type': 'application/json;charset=UTF-8',
    }
    results = []
    for field in ('pn', 'apno'):
        payload = {field: query, 'limit': 20, 'offset': 0}
        if authority:
            payload['authority'] = authority
        body = _do_patent_search(headers, payload)
        results.extend(_extract_patent_list(body, query))
    seen = set()
    unique = []
    for r in results:
        key = r.get('pn', '')
        if key and key not in seen:
            seen.add(key)
            unique.append(r)
    return unique


def _do_patent_search(headers, payload):
    """执行一次专利号搜索请求，带重试。"""
    last_err = None
    for attempt in range(1, MAX_RETRIES + 1):
        try:
            resp = requests.post(PATENT_SEARCH_URL, headers=headers, json=payload, timeout=60)
            if resp.status_code != 200:
                raise RuntimeError(f'HTTP {resp.status_code}: {resp.text[:500]}')
            body = resp.json()
            if isinstance(body, dict) and body.get('error_code') and body.get('error_code') != 0:
                return {'data': {'results': []}}
            return body
        except Exception as e:
            last_err = e
            if attempt < MAX_RETRIES:
                time.sleep(RETRY_INTERVAL)
            else:
                raise RuntimeError(f'专利号查找请求{MAX_RETRIES}次重试后仍失败: {last_err}')


def _extract_patent_list(body, query):
    """从 API 响应中提取专利列表，统一格式。"""
    if not body or not isinstance(body, dict):
        return []
    data = body.get('data') or body
    results = None
    if isinstance(data, dict):
        results = data.get('results') or data.get('list') or data.get('patents')
    if not results:
        results = body.get('results') or body.get('list') or body.get('patents')
    if not results or not isinstance(results, list):
        return []
    extracted = []
    for item in results:
        if not isinstance(item, dict):
            continue
        pn = item.get('pn') or item.get('patent_number') or item.get('publication_number') or ''
        patent_id = item.get('patent_id') or item.get('id') or ''
        title = item.get('title') or item.get('patent_title') or ''
        applicant = item.get('applicant') or item.get('assignee') or ''
        apno = item.get('apno') or item.get('application_number') or ''
        authority_code = item.get('authority') or item.get('country') or item.get('office') or ''
        if pn:
            extracted.append({
                'pn': pn,
                'patent_id': str(patent_id) if patent_id else '',
                'title': title,
                'applicant': applicant,
                'apno': apno,
                'authority': authority_code,
            })
    return extracted


# ── 序列提取 API ──────────────────────────────────────────

def fetch_page(api_key, patent_number, patent_id, extract_type,
               sequence_no_range, sequence_type, sequence_length, page):
    """调用序列提取 API 获取单页结果。"""
    headers = {
        'Authorization': f'Bearer {api_key}',
        'Content-Type': 'application/json;charset=UTF-8',
    }
    payload = {
        'extract_type': extract_type,
        'page': page,
        'size': PAGE_SIZE,
        'sequence_no_range': sequence_no_range,
    }
    if patent_id:
        payload['patent_id'] = patent_id
    elif patent_number:
        payload['patent_number'] = patent_number
    else:
        raise ValueError('patent_number 和 patent_id 至少提供一个')
    if sequence_type:
        payload['sequence_type'] = sequence_type
    if sequence_length:
        payload['sequence_length'] = sequence_length

    last_err = None
    for attempt in range(1, MAX_RETRIES + 1):
        try:
            resp = requests.post(API_URL, headers=headers, json=payload, timeout=120)
            if resp.status_code != 200:
                raise RuntimeError(f'HTTP {resp.status_code}: {resp.text[:500]}')
            body = resp.json()
            if not body.get('status', False):
                err_code = body.get('error_code', 'unknown')
                err_msg = body.get('error_msg') or body.get('error_message') or ''
                raise RuntimeError(f'API错误 error_code={err_code}, error_msg={err_msg}')
            return body
        except Exception as e:
            last_err = e
            if attempt < MAX_RETRIES:
                time.sleep(RETRY_INTERVAL)
            else:
                raise RuntimeError(f'API请求{MAX_RETRIES}次重试后仍失败: {last_err}')


def fetch_all(api_key, patent_number, patent_id, extract_type,
              sequence_no_range, sequence_type, sequence_length):
    """分页全量获取序列。"""
    all_results = []
    total = None
    page = 1
    while True:
        body = fetch_page(api_key, patent_number, patent_id, extract_type,
                          sequence_no_range, sequence_type, sequence_length, page=page)
        data = body.get('data') or {}
        if total is None:
            total = data.get('total', 0)
        results = data.get('results') or []
        all_results.extend(results)
        if not results or len(all_results) >= total:
            break
        page += 1
    return all_results, total


# ── NDJSON 即时落盘 ──────────────────────────────────────

def append_ndjson(filepath, record):
    """将单条记录 append 到 NDJSON 文件（即时落盘）。"""
    with open(filepath, 'a', encoding='utf-8') as f:
        f.write(json.dumps(record, ensure_ascii=False) + '\n')


def load_batch_state(state_path):
    """加载断点恢复状态。"""
    if state_path.exists():
        try:
            with open(state_path, 'r', encoding='utf-8') as f:
                return json.load(f)
        except (json.JSONDecodeError, IOError):
            return {'completed': [], 'failed': [], 'errors': []}
    return {'completed': [], 'failed': [], 'errors': []}


def save_batch_state(state_path, state):
    """保存断点恢复状态。"""
    with open(state_path, 'w', encoding='utf-8') as f:
        json.dump(state, f, ensure_ascii=False, indent=2)


# ── 导出函数 ──────────────────────────────────────────────

def export_excel(sequences, output_path, patent_label):
    """导出 Excel 文件（Summary + Sequences 两个 Sheet）。"""
    if not HAS_OPENPYXL:
        raise RuntimeError('openpyxl 未安装，无法导出 Excel')
    wb = Workbook()
    ws_summary = wb.active
    ws_summary.title = 'Summary'
    type_dist = {}
    for s in sequences:
        st = s.get('sequence_type', 'UNKNOWN')
        type_dist[st] = type_dist.get(st, 0) + 1
    summary_rows = [
        ('专利号', patent_label),
        ('序列总数', len(sequences)),
        ('导出时间', datetime.now().strftime('%Y-%m-%d %H:%M:%S')),
        ('', ''),
        ('序列类型分布', ''),
    ]
    for st, cnt in sorted(type_dist.items()):
        summary_rows.append((st, cnt))
    s_font = Font(bold=True, color='FFFFFF')
    s_fill = PatternFill(start_color='4472C4', end_color='4472C4', fill_type='solid')
    for r, (k, v) in enumerate(summary_rows, 1):
        c1 = ws_summary.cell(row=r, column=1, value=k)
        c2 = ws_summary.cell(row=r, column=2, value=v)
        if r == 1 or (k and not v and r == 5):
            c1.font = s_font
            c1.fill = s_fill
    ws_summary.column_dimensions['A'].width = 20
    ws_summary.column_dimensions['B'].width = 35

    ws = wb.create_sheet('Sequences')
    loc_keys = set()
    for s in sequences:
        loc = s.get('location') or {}
        loc_keys.update(loc.keys())
    loc_keys = sorted(loc_keys)
    headers = (
        ['Seq ID', 'Sequence ID Number', 'Sequence Type', 'Length']
        + [f'Location_{k}' for k in loc_keys]
        + ['Sequence']
    )
    h_font = Font(bold=True, color='FFFFFF')
    h_fill = PatternFill(start_color='4472C4', end_color='4472C4', fill_type='solid')
    h_align = Alignment(horizontal='center', vertical='center', wrap_text=True)
    for col, h in enumerate(headers, 1):
        c = ws.cell(row=1, column=col, value=h)
        c.font = h_font
        c.fill = h_fill
        c.alignment = h_align
    for r, s in enumerate(sequences, 2):
        sid = s.get('seq_id', '')
        sid_num = s.get('sequence_id_number', [])
        sid_num_str = ', '.join(str(x) for x in sid_num) if isinstance(sid_num, list) else str(sid_num)
        stype = s.get('sequence_type', '')
        slen = s.get('s_len', '')
        seq = s.get('s_sequence', '')
        loc = s.get('location') or {}
        col = 1
        ws.cell(row=r, column=col, value=sid); col += 1
        ws.cell(row=r, column=col, value=sid_num_str); col += 1
        ws.cell(row=r, column=col, value=stype); col += 1
        ws.cell(row=r, column=col, value=slen); col += 1
        for lk in loc_keys:
            lv = loc.get(lk, [])
            lv_str = ', '.join(str(x) for x in lv) if isinstance(lv, list) else str(lv)
            ws.cell(row=r, column=col, value=lv_str); col += 1
        ws.cell(row=r, column=col, value=seq)
    ws.freeze_panes = 'A2'
    try:
        wb.save(output_path)
    except PermissionError:
        output_path = _timestamped_path(output_path)
        wb.save(output_path)
    return output_path


def export_fasta(sequences, output_path, patent_label):
    """导出 FASTA 文件。"""
    def write_inner(f):
        for s in sequences:
            sid = s.get('seq_id', '')
            sid_num = s.get('sequence_id_number', [])
            sid_num_str = ','.join(str(x) for x in sid_num) if isinstance(sid_num, list) else str(sid_num)
            stype = s.get('sequence_type', 'UNKNOWN')
            slen = s.get('s_len', 0)
            seq = s.get('s_sequence', '')
            header = f'>{patent_label}|seq_id_no:{sid_num_str}|type:{stype}|len:{slen}|seq_id:{sid}'
            f.write(header + '\n')
            for i in range(0, len(seq), 60):
                f.write(seq[i:i + 60] + '\n')
    try:
        with open(output_path, 'w') as f:
            write_inner(f)
    except PermissionError:
        output_path = _timestamped_path(output_path)
        with open(output_path, 'w') as f:
            write_inner(f)
    return output_path


def export_json(sequences, output_path, patent_label):
    """导出 JSON 文件。"""
    data = {
        'patent': patent_label,
        'total': len(sequences),
        'export_time': datetime.now().strftime('%Y-%m-%d %H:%M:%S'),
        'sequences': sequences,
    }
    with open(output_path, 'w', encoding='utf-8') as f:
        json.dump(data, f, ensure_ascii=False, indent=2)
    return output_path


# ── 批量提取主函数 ────────────────────────────────────────

def run_batch(patents_raw, api_key, seq_id_range='all',
              sequence_type=None, seq_length_start=None, seq_length_end=None,
              output_format='excel', output_dir=None,
              progress_callback=None):
    """
    批量提取多个专利的生物序列。

    参数:
        patents_raw: 专利号字符串（逗号/分号/换行分隔）
        api_key: 智慧芽 API Key
        seq_id_range: 'all' / 具体编号 / 范围
        sequence_type: 'DNA' / 'RNA' / 'Protein' / None
        seq_length_start: 长度下限
        seq_length_end: 长度上限
        output_format: 'excel' / 'fasta' / 'json' / 'both'
        output_dir: 输出目录
        progress_callback: 进度回调函数

    返回:
        batch_summary dict
    """
    # Phase 1: 前置条件检查
    if not api_key:
        raise ValueError('API Key 不能为空')
    patent_list = parse_patent_ids(patents_raw)
    if not patent_list:
        raise ValueError('专利号列表不能为空')

    # 准备输出目录
    if output_dir:
        out_dir = Path(output_dir)
    else:
        out_dir = Path.cwd() / 'outputs'
    out_dir.mkdir(parents=True, exist_ok=True)

    ndjson_path = out_dir / 'sequences_all.ndjson'
    state_path = out_dir / 'batch_state.json'
    summary_path = out_dir / 'batch_summary.json'

    # 加载断点状态
    state = load_batch_state(state_path)
    completed = set(state.get('completed', []))

    # 解析筛选条件
    extract_type, seq_no_range = parse_seq_id_range(seq_id_range)
    seq_len = None
    if seq_length_start is not None or seq_length_end is not None:
        seq_len = {
            'start': int(seq_length_start) if seq_length_start is not None else 0,
            'end': int(seq_length_end) if seq_length_end is not None else 999999,
        }

    # Phase 2: 逐专利提取
    results = {}
    total_patents = len(patent_list)

    for idx, patent_input in enumerate(patent_list):
        if patent_input in completed:
            if progress_callback:
                progress_callback(idx + 1, total_patents, patent_input, 0, 'skip')
            continue

        patent_label = patent_input
        try:
            # 专利号标准化
            if progress_callback:
                progress_callback(idx + 1, total_patents, patent_input, 0, 'standardizing')

            pn_results = search_patent_pn(api_key, patent_input)
            if not pn_results:
                state['failed'].append(patent_input)
                state['errors'].append({
                    'patent': patent_input,
                    'error': '未找到匹配的专利'
                })
                save_batch_state(state_path, state)
                if progress_callback:
                    progress_callback(idx + 1, total_patents, patent_input, 0, 'error')
                continue

            # 非交互式：取第一个结果
            resolved = pn_results[0]
            patent_id = resolved.get('patent_id', '')
            patent_label = resolved.get('pn', patent_input)

            if not patent_id:
                state['failed'].append(patent_input)
                state['errors'].append({
                    'patent': patent_input,
                    'error': '找到专利但未返回 UUID'
                })
                save_batch_state(state_path, state)
                if progress_callback:
                    progress_callback(idx + 1, total_patents, patent_input, 0, 'error')
                continue

            # 序列提取
            if progress_callback:
                progress_callback(idx + 1, total_patents, patent_label, 0, 'extracting')

            sequences, total = fetch_all(
                api_key=api_key,
                patent_number='',
                patent_id=patent_id,
                extract_type=extract_type,
                sequence_no_range=seq_no_range,
                sequence_type=sequence_type,
                sequence_length=seq_len,
            )

            # 即时落盘
            for seq in sequences:
                seq_record = dict(seq)
                seq_record['patent_id'] = patent_id
                seq_record['patent_label'] = patent_label
                append_ndjson(ndjson_path, seq_record)

            results[patent_label] = {
                'patent_input': patent_input,
                'patent_id': patent_id,
                'sequences': sequences,
                'total': total,
                'status': 'success'
            }

            state['completed'].append(patent_input)
            save_batch_state(state_path, state)

            if progress_callback:
                progress_callback(idx + 1, total_patents, patent_label, len(sequences), 'done')

        except Exception as e:
            state['failed'].append(patent_input)
            state['errors'].append({
                'patent': patent_input,
                'error': str(e)
            })
            save_batch_state(state_path, state)
            if progress_callback:
                progress_callback(idx + 1, total_patents, patent_input, 0, 'error')

        # 专利间间隔（避免限流）
        if idx < total_patents - 1:
            time.sleep(2)

    # Phase 3: 合并导出
    if progress_callback:
        progress_callback(0, total_patents, '', 0, 'exporting')

    files = []
    for patent_label, info in results.items():
        seqs = info['sequences']
        if output_format in ('excel', 'both') and HAS_OPENPYXL:
            p = str(out_dir / f'{patent_label}_sequences.xlsx')
            actual = export_excel(seqs, p, patent_label)
            files.append(Path(actual).name)
        if output_format in ('fasta', 'both'):
            p = str(out_dir / f'{patent_label}_sequences.fasta')
            actual = export_fasta(seqs, p, patent_label)
            files.append(Path(actual).name)
        if output_format == 'json':
            p = str(out_dir / f'{patent_label}_sequences.json')
            actual = export_json(seqs, p, patent_label)
            files.append(Path(actual).name)

    # 生成 batch_summary
    type_dist_all = {}
    for info in results.values():
        for s in info['sequences']:
            st = s.get('sequence_type', 'UNKNOWN')
            type_dist_all[st] = type_dist_all.get(st, 0) + 1

    summary = {
        'export_time': datetime.now().strftime('%Y-%m-%d %H:%M:%S'),
        'total_patents': total_patents,
        'succeeded': len(results),
        'failed': len(state.get('failed', [])),
        'total_sequences': sum(len(info['sequences']) for info in results.values()),
        'type_distribution': type_dist_all,
        'files': files,
        'per_patent': {
            label: {
                'total': info['total'],
                'fetched': len(info['sequences']),
                'status': info['status'],
            }
            for label, info in results.items()
        },
        'errors': state.get('errors', []),
        'ndjson_path': str(ndjson_path),
        'output_dir': str(out_dir),
    }

    with open(summary_path, 'w', encoding='utf-8') as f:
        json.dump(summary, f, ensure_ascii=False, indent=2)

    if progress_callback:
        progress_callback(0, total_patents, '', summary['total_sequences'], 'complete')

    return summary


# ── CLI 入口 ──────────────────────────────────────────────

def main():
    parser = argparse.ArgumentParser(
        description='专利生物序列批量提取工具 v1.0',
        formatter_class=argparse.RawDescriptionHelpFormatter,
        epilog='''
示例:
  python batch_extract.py --patents WO2020123456 --api-key sk-xxx --seq-range all
  python batch_extract.py --patents WO1,WO2,WO3 --api-key sk-xxx --seq-type Protein
  python batch_extract.py --patents WO2020123456 --api-key-file key.txt --seq-range 1-50 --format fasta
        ''')
    parser.add_argument('--patents', required=True, help='专利号列表，逗号/分号/换行分隔')
    parser.add_argument('--api-key', default=None, help='智慧芽 Bio OpenAPI Key (sk-xxx)')
    parser.add_argument('--api-key-file', default=None, help='从文件读取 API Key')
    parser.add_argument('--seq-range', default='all', help='Seq ID 范围: all / 具体编号 / 范围(如 1-50)')
    parser.add_argument('--seq-type', default=None, choices=['DNA', 'RNA', 'Protein', None],
                        help='序列类型筛选')
    parser.add_argument('--seq-len-start', type=int, default=None, help='序列长度下限')
    parser.add_argument('--seq-len-end', type=int, default=None, help='序列长度上限')
    parser.add_argument('--format', default='excel', choices=['excel', 'fasta', 'json', 'both'],
                        help='输出格式')
    parser.add_argument('--output-dir', default=None, help='输出目录（默认 ./outputs/）')
    args = parser.parse_args()

    # 环境变量兜底
    api_key = args.api_key or os.environ.get('ZHIHUIYA_API_KEY', '')
    if not api_key and args.api_key_file:
        api_key = read_api_key('', args.api_key_file)
    if not api_key:
        print('[ERROR] 未提供 API Key，请使用 --api-key 或 --api-key-file 或设置 ZHIHUIYA_API_KEY 环境变量')
        sys.exit(1)

    def cli_progress(current, total, patent, count, phase):
        if phase == 'standardizing':
            print(f'[{current}/{total}] 标准化专利号: {patent} ...')
        elif phase == 'extracting':
            print(f'[{current}/{total}] 提取序列: {patent} ...')
        elif phase == 'done':
            print(f'[{current}/{total}] ✅ {patent} → {count} 条序列')
        elif phase == 'skip':
            print(f'[{current}/{total}] ⏭️ 跳过已完成: {patent}')
        elif phase == 'error':
            print(f'[{current}/{total}] ❌ 失败: {patent}')
        elif phase == 'exporting':
            print('正在导出文件 ...')
        elif phase == 'complete':
            print(f'✅ 全部完成，共 {count} 条序列')

    try:
        summary = run_batch(
            patents_raw=args.patents,
            api_key=api_key,
            seq_id_range=args.seq_range,
            sequence_type=args.seq_type,
            seq_length_start=args.seq_len_start,
            seq_length_end=args.seq_len_end,
            output_format=args.format,
            output_dir=args.output_dir,
            progress_callback=cli_progress,
        )
        print(f'\n=== 批量提取汇总 ===')
        print(f'专利数: {summary["succeeded"]}/{summary["total_patents"]} 成功')
        print(f'序列总数: {summary["total_sequences"]}')
        print(f'输出目录: {summary["output_dir"]}')
        if summary['files']:
            print(f'输出文件: {", ".join(summary["files"])}')
        if summary['errors']:
            print(f'失败专利: {len(summary["errors"])} 个')
            for err in summary['errors']:
                print(f'  - {err["patent"]}: {err["error"]}')
    except ValueError as e:
        print(f'[ERROR] {e}')
        sys.exit(1)
    except Exception as e:
        print(f'[ERROR] {e}')
        sys.exit(1)


if __name__ == '__main__':
    main()
